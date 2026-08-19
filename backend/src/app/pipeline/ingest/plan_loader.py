"""Reads the CMMS workbook and the contractor CSV.

Two small feeds with one thing in common: neither states a timezone, and each references
equipment by something other than the plant's own asset tag.
"""

import csv
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.pipeline.ingest.counters import IngestCounters
from app.pipeline.ingest.timestamps import from_excel_serial
from app.pipeline.manifest import CONTRACTOR_FILE, MAINTENANCE_FILE, SOURCES

_MAINTENANCE_PROVENANCE = (MAINTENANCE_FILE, SOURCES[MAINTENANCE_FILE]["department"])
_CONTRACTOR_PROVENANCE = (CONTRACTOR_FILE, SOURCES[CONTRACTOR_FILE]["department"])


def load_work_orders(path: Path, counters: IngestCounters) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["WorkOrders"]
        rows = []
        for record in sheet.iter_rows(min_row=2, values_only=True):
            row = _work_order(record, counters)
            if row is not None:
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _work_order(record: tuple[Any, ...], counters: IngestCounters) -> tuple[Any, ...] | None:
    work_order, historian_tag, wo_type, priority, plant, serial, raised_by = record
    try:
        scheduled = from_excel_serial(float(serial))
    except (TypeError, ValueError):
        counters.rejected(MAINTENANCE_FILE)
        return None

    counters.timestamp(assumed=scheduled.tz_assumed)
    counters.kept(MAINTENANCE_FILE)
    return (
        work_order,
        historian_tag,
        wo_type,
        priority,
        plant,
        scheduled.at,
        scheduled.tz_assumed,
        raised_by,
        *_MAINTENANCE_PROVENANCE,
    )


def load_callouts(path: Path, counters: IngestCounters) -> list[tuple[Any, ...]]:
    rows = []
    with path.open(encoding="utf-8", newline="") as handle:
        for record in csv.DictReader(handle):
            row = _callout(record, counters)
            if row is not None:
                rows.append(row)
    return rows


def _callout(record: dict[str, str], counters: IngestCounters) -> tuple[Any, ...] | None:
    try:
        serviced_on = date.fromisoformat(record["service_date"])
        hours = float(record["hours"])
    except (KeyError, ValueError):
        counters.rejected(CONTRACTOR_FILE)
        return None

    counters.kept(CONTRACTOR_FILE)
    return (
        record["callout_id"],
        record["vendor"],
        record["equipment_ref"],
        serviced_on,
        record["billing_status"],
        hours,
        *_CONTRACTOR_PROVENANCE,
    )
